"""
生成「生存竞技场 · 游戏介绍」通用版作品集 PPT。
深色主题，与 HTML 版视觉一致。
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path
from PIL import Image

# ── 主题色 ──
BG      = RGBColor(0x0A, 0x0A, 0x0F)
CARD    = RGBColor(0x16, 0x1A, 0x26)
CARD2   = RGBColor(0x1C, 0x21, 0x30)
LINE    = RGBColor(0x2A, 0x31, 0x44)
TEXT    = RGBColor(0xE8, 0xEB, 0xF2)
MUTED   = RGBColor(0x98, 0xA2, 0xB8)
DIM     = RGBColor(0x6B, 0x74, 0x88)
ACCENT  = RGBColor(0xE9, 0x45, 0x60)
ACCENT2 = RGBColor(0xFF, 0x6B, 0x8A)
GOLD    = RGBColor(0xF0, 0xC0, 0x40)
GEM     = RGBColor(0xA8, 0x55, 0xF7)
GREEN   = RGBColor(0x34, 0xD3, 0x99)
AMBER   = RGBColor(0xFB, 0xBF, 0x24)
BLUE    = RGBColor(0x60, 0xA5, 0xFA)

FONT = "微软雅黑"
BASE = Path(__file__).parent
ASSETS = BASE / "assets"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def bg(slide, color=BG):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background(); s.shadow.inherit = False
    return s


def box(slide, l, t, w, h, fill=CARD, line=None, radius=None):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape, l, t, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line; s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if radius:
        try: s.adjustments[0] = radius
        except Exception: pass
    return s


def text(slide, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=1.0, space_after=0):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [[(runs, 14, TEXT, False)]]
    first = True
    for line_runs in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align; p.line_spacing = spacing; p.space_after = Pt(space_after)
        for item in line_runs:
            txt, size, color, bold = item[0], item[1], item[2], item[3]
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.color.rgb = color
            r.font.bold = bold; r.font.name = FONT
    return tb


def accent_bar(slide, l, t, w=Inches(0.06), h=Inches(0.42), color=ACCENT):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background(); s.shadow.inherit = False
    return s


def header(slide, num, title, sub=None):
    accent_bar(slide, Inches(0.7), Inches(0.62))
    text(slide, Inches(0.92), Inches(0.58), Inches(11), Inches(0.3), [[(num, 12, ACCENT, True)]])
    text(slide, Inches(0.92), Inches(0.85), Inches(11.9), Inches(0.5), [[(title, 27, TEXT, True)]])
    y = Inches(1.42)
    if sub:
        text(slide, Inches(0.92), Inches(1.36), Inches(11.6), Inches(0.5),
             [[(sub, 12.5, MUTED, False)]], spacing=1.22)
        y = Inches(1.9)
    return y


def pic_fit(slide, path, l, t, max_w, max_h):
    with Image.open(path) as im:
        iw, ih = im.size
    ratio = min(max_w / iw, max_h / ih)
    w, h = int(iw * ratio), int(ih * ratio)
    left = l + int((max_w - w) / 2)
    top = t + int((max_h - h) / 2)
    return slide.shapes.add_picture(str(path), left, top, w, h)


def hexc(s):
    s = s.lstrip("#")
    return RGBColor(int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))


def table(slide, l, t, w, col_w, heads, rows, row_h=Inches(0.5), head_h=Inches(0.42),
          font_size=11, colors=None):
    box(slide, l, t, w, head_h, fill=CARD2)
    cx = l + Inches(0.18)
    for h_, cw in zip(heads, col_w):
        text(slide, cx, t, Inches(cw), head_h, [[(h_, 10.5, DIM, True)]], anchor=MSO_ANCHOR.MIDDLE)
        cx += Inches(cw)
    y = t + head_h
    for i, row in enumerate(rows):
        box(slide, l, y, w, row_h, fill=(CARD if i % 2 == 0 else BG))
        cx = l + Inches(0.18)
        for j, (val, cw) in enumerate(zip(row, col_w)):
            col = TEXT if j == 0 else MUTED
            bold = (j == 0)
            if colors and colors.get(j):
                col = colors[j]
            text(slide, cx, y, Inches(cw), row_h, [[(val, font_size, col, bold)]],
                 anchor=MSO_ANCHOR.MIDDLE)
            cx += Inches(cw)
        y += row_h
    return y


# ══════════════════════════════════════════════
# 1. 封面
# ══════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); bg(s)
box(s, 0, 0, SW, Inches(0.05), fill=ACCENT)
text(s, Inches(1.0), Inches(1.3), Inches(6), Inches(0.3),
     [[("浏览器即玩 · 无需下载 · 手机电脑双端", 13, ACCENT2, True)]])
text(s, Inches(1.0), Inches(1.75), Inches(11.3), Inches(2.0),
     [[("🔥 生存竞技场", 50, TEXT, True)],
      [("Survival Arena", 30, ACCENT2, True)]],
     spacing=1.1, space_after=8)
text(s, Inches(1.0), Inches(3.8), Inches(9.9), Inches(1.1),
     [[("一款用原生 JavaScript + Canvas 从零手写的 2D 俯视角波次生存射击游戏。", 14, MUTED, False)],
      [("不用下载、不用安装，打开浏览器就能玩——手机和电脑都能玩，", 14, MUTED, False)],
      [("死了还能让好友把你救回来。", 14, MUTED, False)]],
     spacing=1.3)
chips = ["浏览器 / 手机 / 电脑", "波次生存射击",
         "4 武器 · 5 敌人 · 6 皮肤 · 5 段位", "已上线可玩"]
x = Inches(1.0)
for c in chips:
    w = Inches(0.42) + Emu(int(len(c) * 118000))
    box(s, x, Inches(5.35), w, Inches(0.42), fill=CARD, line=LINE, radius=0.35)
    text(s, x, Inches(5.35), w, Inches(0.42), [[(c, 11, MUTED, False)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    x += w + Inches(0.14)
text(s, Inches(1.0), Inches(6.2), Inches(11), Inches(0.35),
     [[("▶ 立刻开始游玩：https://survival-arena.com/", 13, GREEN, True)]])

# ══════════════════════════════════════════════
# 2. 这是个什么游戏
# ══════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); bg(s)
header(s, "01 / 这是个什么游戏", "一句话说清：点开就打的爽快小游戏",
       "没有下载、没有安装、没有冗长注册。打开网页，WASD 走动、鼠标瞄准射击，一波波敌人冲上来，撑得越久越强。")

stats = [("0 安装", "浏览器打开即玩", "手机电脑通用"),
         ("4 种", "武器", "手枪/霰弹/冲锋/狙击"),
         ("5 种", "敌人", "每 5 波一个 BOSS"),
         ("6 款", "皮肤", "4 个稀有度梯度"),
         ("12 项", "成就", "另有每日挑战"),
         ("4 种", "复活方式", "含跨设备好友复活")]
cw = Inches(1.83); gap = Inches(0.14); x0 = Inches(0.92)
for i, (num, lab, sub) in enumerate(stats):
    col, row = i % 3, i // 3
    x = x0 + col * (cw + gap)
    y = Inches(1.95) + row * Inches(1.15)
    box(s, x, y, cw, Inches(1.0), fill=CARD, line=LINE, radius=0.06)
    box(s, x, y, cw, Inches(0.045), fill=ACCENT)
    text(s, x + Inches(0.22), y + Inches(0.16), cw - Inches(0.4), Inches(0.4),
         [[(num, 23, TEXT, True)]])
    text(s, x + Inches(0.22), y + Inches(0.55), cw - Inches(0.4), Inches(0.22),
         [[(lab, 10.5, MUTED, False)]])
    text(s, x + Inches(0.22), y + Inches(0.75), cw - Inches(0.4), Inches(0.22),
         [[(sub, 9, DIM, False)]])

box(s, Inches(6.82), Inches(1.95), Inches(5.6), Inches(2.15), fill=CARD2, line=LINE, radius=0.03)
pic_fit(s, ASSETS / "02-menu.png", Inches(6.95), Inches(2.05), Inches(5.34), Inches(1.95))
text(s, Inches(0.92), Inches(4.35), Inches(11.5), Inches(2.4),
     [[("主菜单：所有功能收在一个界面", 14, ACCENT2, True)],
      [("开始战斗、装备升级、每日奖励、排行榜、皮肤——进来就知道该点哪。", 12, MUTED, False)]],
     spacing=1.3)
box(s, Inches(0.92), Inches(5.0), Inches(5.6), Inches(1.6), fill=CARD2, line=LINE, radius=0.03)
pic_fit(s, ASSETS / "01-login.png", Inches(1.02), Inches(5.1), Inches(5.4), Inches(1.4))
box(s, Inches(6.82), Inches(5.0), Inches(5.6), Inches(1.6), fill=CARD2, line=LINE, radius=0.03)
pic_fit(s, ASSETS / "dashboard.png", Inches(6.92), Inches(5.1), Inches(5.4), Inches(1.4))
text(s, Inches(0.92), Inches(6.7), Inches(11.5), Inches(0.3),
     [[("左：登录界面（首次登录自动注册）　　右：线上数据看板（实时读取埋点数据）", 10, DIM, False)]])

# ══════════════════════════════════════════════
# 3. 怎么玩：五步循环
# ══════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); bg(s)
header(s, "02 / 怎么玩", "从打开到变强，五步走完",
       "整个游戏循环：进去 → 打 → 变强 → 再来一局。每一步都有明确的正反馈。")

steps = [("1", "登录进入", "手机号验证码或邮箱登录，首次登录自动注册，几秒进游戏；数据云端同步，换设备接着玩。"),
         ("2", "开始战斗", "电脑 WASD + 鼠标瞄准，空格冲刺；手机虚拟摇杆 + 射击键，滑动闪避。第一波敌人很温和。"),
         ("3", "击杀敌人收集星星", "击杀掉落 ⭐ 星星，波次清空也有奖励；星星用来买装备、升武器，连杀有额外奖励。"),
         ("4", "每 5 波打 BOSS", "敌人越来越强；第 5/10/15/20 波是 BOSS 战，血厚 + 远程弹幕 + 带护卫小兵。"),
         ("5", "阵亡 → 复活 → 变强", "可看广告/用令牌/花钻石/邀好友复活；金币钻石跨局累积，永久提升属性，下局更强。")]
y = Inches(1.9)
for n, title, desc in steps:
    box(s, Inches(0.92), y, Inches(11.5), Inches(0.95), fill=CARD, line=LINE, radius=0.05)
    box(s, Inches(0.92), y, Inches(0.05), Inches(0.95), fill=ACCENT)
    box(s, Inches(1.15), y + Inches(0.26), Inches(0.42), Inches(0.42),
        fill=CARD2, line=ACCENT, radius=0.2)
    text(s, Inches(1.15), y + Inches(0.26), Inches(0.42), Inches(0.42),
         [[(n, 15, ACCENT2, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(1.78), y + Inches(0.14), Inches(2.6), Inches(0.3), [[(title, 13.5, TEXT, True)]])
    text(s, Inches(1.78), y + Inches(0.47), Inches(10.4), Inches(0.42), [[(desc, 10.5, MUTED, False)]])
    y += Inches(1.06)

# ══════════════════════════════════════════════
# 4. 战斗实况
# ══════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); bg(s)
header(s, "02 / 怎么玩", "真实战斗画面")
for i, (fn, title, cap) in enumerate([
    ("07-battle.png", "开局", "手枪起步，敌人从四周涌来"),
    ("10-combat.png", "混战", "屏幕震动 + 粒子特效，越打越密集"),
]):
    l = Inches(0.92) + i * Inches(5.95)
    box(s, l, Inches(1.85), Inches(5.6), Inches(4.75), fill=CARD2, line=LINE, radius=0.03)
    pic_fit(s, ASSETS / fn, l + Inches(0.12), Inches(1.97), Inches(5.36), Inches(3.95))
    text(s, l + Inches(0.2), Inches(6.03), Inches(5.2), Inches(0.4),
         [[(title + "　", 12.5, TEXT, True), (cap, 10.5, MUTED, False)]], spacing=1.2)

# ══════════════════════════════════════════════
# 5. 核心功能总览
# ══════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); bg(s)
header(s, "03 / 核心功能", "六个撑起整个游戏的系统",
       "每个功能都解决一个具体问题，而不是「为了有而有」")

feats = [("🔫 武器系统", "四种枪，各有一套升级树",
          "手枪免费起手，霰弹枪 200⭐ / 冲锋枪 350⭐ / 狙击枪 500⭐；解锁后永久拥有、跨局保留，每把 3 条升级路线各 10 级。"),
         ("💰 三货币经济", "每种钱花在不同地方",
          "⭐星星管这一局（买装备升武器）、🪙金币管长期（永久属性）、💎钻石管外观与便利（皮肤/复活）。"),
         ("🎨 皮肤系统", "6 款皮肤，4 个稀有度",
          "普通/稀有/史诗/传说四档，价格 60–200💎；全部 Canvas 程序化绘制，商店卡片是实时动效预览。"),
         ("🤝 好友复活", "死了发个链接，朋友把你捞回来",
          "阵亡生成专属链接，好友点开即满血复活 + 3 秒无敌；跨设备互通，双方都有奖励。"),
         ("🏆 成长与竞争", "段位、成就、每日挑战",
          "5 大段位按最高波次解锁（青铜→传说）；12 项成就 + 每日挑战 + 7 天连签。"),
         ("📱 手机完整适配", "不是「能打开」，是「玩得爽」",
          "虚拟摇杆 + 触控瞄准 + 滑动手势闪避 + 刘海屏安全区 + 性能自动降级保帧率。")]
y0 = Inches(1.9)
cw2, gap2 = Inches(5.6), Inches(0.3)
for i, (title, sub, desc) in enumerate(feats):
    col, row = i % 2, i // 2
    x = Inches(0.92) + col * (cw2 + gap2)
    yy = y0 + row * Inches(2.55)
    box(s, x, yy, cw2, Inches(2.35), fill=CARD, line=LINE, radius=0.05)
    box(s, x, yy, Inches(0.05), Inches(2.35), fill=ACCENT)
    text(s, x + Inches(0.28), yy + Inches(0.2), Inches(5.0), Inches(0.3),
         [[(title, 15, TEXT, True)]])
    text(s, x + Inches(0.28), yy + Inches(0.56), Inches(5.0), Inches(0.28),
         [[(sub, 11.5, ACCENT2, True)]])
    text(s, x + Inches(0.28), yy + Inches(0.92), Inches(5.05), Inches(1.3),
         [[(desc, 10.5, MUTED, False)]], spacing=1.3)

# ══════════════════════════════════════════════
# 6. 武器系统
# ══════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); bg(s)
header(s, "03 / 核心功能", "武器系统：四种枪，各有一套升级树",
       "解锁后永久拥有、跨局保留；每把枪 3 条独立升级路线（伤害 / 射速 / 特殊），最高各 10 级")

rows = [("🔫 手枪", "免费", "平衡型，1 发弹丸", "初始"),
        ("💥 霰弹枪", "200 ⭐", "5 发扇形散射，近距离威力巨大", "2–3 局"),
        ("⚡ 冲锋枪", "350 ⭐", "极高射速，弹幕压制", "3–4 局"),
        ("🎯 狙击枪", "500 ⭐", "3 倍伤害，一击制敌", "5–6 局")]
table(s, Inches(0.92), Inches(2.0), Inches(11.5), [2.4, 2.0, 5.0, 2.1],
      ["武器", "解锁费用", "特点", "约需局数"], rows, row_h=Inches(0.72), font_size=11.5)

box(s, Inches(0.92), Inches(5.25), Inches(11.5), Inches(1.6), fill=CARD, line=GEM, radius=0.05)
text(s, Inches(1.25), Inches(5.42), Inches(10.9), Inches(1.3),
     [[("设计意图", 13, GEM, True)],
      [("星星是局内货币，买到的武器却跨局保留——这让「这一局打得不好」不会白费，", 11.5, MUTED, False)],
      [("每一局都在为解锁下一把枪积累进度。四个价格档位（免费 / 2–3 局 / 3–4 局 / 5–6 局）", 11.5, MUTED, False)],
      [("刚好构成一条「越玩越有」的成长节奏。", 11.5, MUTED, False)]],
     spacing=1.3, space_after=2)

# ══════════════════════════════════════════════
# 7. 三货币经济
# ══════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); bg(s)
header(s, "03 / 核心功能", "三货币经济：每种钱花在不同的地方")

rows = [("⭐ 星星", "局内", "击杀敌人、清波奖励、连杀", "买装备、升武器（跨局保留）"),
        ("🪙 金币", "跨局累积", "签到、每日挑战、钻石兑换", "永久提升角色属性、买自动射击"),
        ("💎 钻石", "跨局累积", "充值、签到奖励、成就", "买皮肤、复活、兑换金币/星星")]
table(s, Inches(0.92), Inches(1.95), Inches(11.5), [2.0, 2.0, 3.8, 3.7],
      ["货币", "范围", "怎么来", "用来做什么"], rows, row_h=Inches(0.78), font_size=11)

box(s, Inches(0.92), Inches(4.85), Inches(11.5), Inches(1.15), fill=CARD, line=GREEN, radius=0.05)
text(s, Inches(1.25), Inches(5.0), Inches(10.9), Inches(0.9),
     [[("设计意图", 12.5, GREEN, True),
       ("：三种货币各管一段——星星管「这一局怎么打」，金币管「长期怎么变强」，", 11, MUTED, False)],
      [("钻石管「外观和便利」。这样玩家每一局都有收获，又不会因为一局打崩就白玩。", 11, MUTED, False)]],
     spacing=1.3, space_after=2)

box(s, Inches(0.92), Inches(6.15), Inches(11.5), Inches(0.75), fill=CARD2, line=LINE, radius=0.06)
text(s, Inches(1.25), Inches(6.15), Inches(10.9), Inches(0.75),
     [[("常用道具：", 11.5, ACCENT2, True),
       ("生命恢复 30⭐ ｜ 护盾 45⭐ ｜ 速度提升 30⭐ ｜ 双倍伤害 60⭐ ｜ 星星磁铁 40⭐", 11, MUTED, False)]],
     anchor=MSO_ANCHOR.MIDDLE)

# ══════════════════════════════════════════════
# 8. 皮肤系统
# ══════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); bg(s)
header(s, "03 / 核心功能", "皮肤系统：6 款皮肤，4 个稀有度",
       "全部用 Canvas 程序化绘制（零美术图片资源），商店卡片是实时渲染的动图，不是静态图")

skins = [("skin-1-default.png", "制式战士", "普通", "#8FA3B8", "免费 / 静态"),
         ("skin-2-flame.png", "烈焰使者", "稀有", "#4FC3F7", "60💎 / 呼吸旋转"),
         ("skin-3-shadow.png", "暗影刺客", "稀有", "#4FC3F7", "60💎 / 残影"),
         ("skin-4-neon.png", "霓虹战士", "史诗", "#C77DFF", "90💎 / 发光扫描"),
         ("skin-5-gold.png", "黄金骑士", "史诗", "#C77DFF", "120💎 / 旋转高光"),
         ("skin-6-void.png", "虚空领主", "传说", "#FFB703", "200💎 / 全套动效")]
cw = Inches(1.83); gap = Inches(0.14); x0 = Inches(0.92)
for i, (fn, name, rar, color, meta) in enumerate(skins):
    x = x0 + i * (cw + gap)
    box(s, x, Inches(1.95), cw, Inches(3.15), fill=CARD, line=LINE, radius=0.05)
    box(s, x, Inches(1.95), cw, Inches(0.05), fill=hexc(color))
    pic_fit(s, ASSETS / fn, x + Inches(0.12), Inches(2.12), cw - Inches(0.24), Inches(1.95))
    text(s, x + Inches(0.14), Inches(4.2), cw - Inches(0.28), Inches(0.28), [[(name, 12.5, TEXT, True)]])
    text(s, x + Inches(0.14), Inches(4.5), cw - Inches(0.28), Inches(0.24), [[(rar, 10, hexc(color), True)]])
    text(s, x + Inches(0.14), Inches(4.74), cw - Inches(0.28), Inches(0.24), [[(meta, 9, DIM, False)]])

box(s, Inches(0.92), Inches(5.3), Inches(11.5), Inches(1.55), fill=CARD, line=LINE, radius=0.04)
text(s, Inches(1.25), Inches(5.48), Inches(10.9), Inches(1.25),
     [[("为什么这么设计", 13, ACCENT2, True)],
      [("价格严格对齐「动效密度」——多花一档的钱，确实多买到一层动态（稀有度越高，动效越多）。", 11, MUTED, False)],
      [("这既让加价变得可解释，也给了玩家升级的理由。另外每套皮肤都先保证「一眼能认出自己」，再叠加主题装饰。", 11, MUTED, False)]],
     spacing=1.3, space_after=2)

# ══════════════════════════════════════════════
# 9. 好友复活（截图 + 说明）
# ══════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); bg(s)
header(s, "03 / 核心功能", "好友复活：死了发个链接，朋友把你捞回来",
       "游戏最有意思的机制 —— 把「死亡」从挫败点变成了社交点")

box(s, Inches(0.92), Inches(1.95), Inches(5.4), Inches(4.7), fill=CARD2, line=LINE, radius=0.03)
pic_fit(s, ASSETS / "11-revive.png", Inches(1.03), Inches(2.05), Inches(5.18), Inches(4.5))

text(s, Inches(6.6), Inches(2.0), Inches(5.8), Inches(0.4),
     [[("怎么运作", 14, ACCENT2, True)]])
text(s, Inches(6.6), Inches(2.45), Inches(5.8), Inches(2.0),
     [[("① 阵亡后生成一条专属分享链接（含唯一标识）", 11.5, MUTED, False)],
      [("② 发给好友，好友点开进游戏 → 云端自动记录", 11.5, MUTED, False)],
      [("③ 核验通过 → 你满血复活 + 3 秒无敌护盾", 11.5, MUTED, False)],
      [("④ 双方都有奖励：各 +50 🪙 金币 +5 💎 钻石", 11.5, MUTED, False)],
      [("", 8, MUTED, False)],
      [("跨设备互通：好友在电脑、你在手机，照样救得回来。", 11.5, GREEN, True)]],
     spacing=1.32, space_after=3)

text(s, Inches(6.6), Inches(4.6), Inches(5.8), Inches(0.4),
     [[("三种防刷设计", 14, ACCENT2, True)]])
text(s, Inches(6.6), Inches(5.05), Inches(5.8), Inches(1.6),
     [[("· 设备 ID：同一浏览器设备不重复计数", 11.5, MUTED, False)],
      [("· 账号哈希：同账号跨设备也算同一人", 11.5, MUTED, False)],
      [("· 本地历史：记录已使用的复活来源", 11.5, MUTED, False)],
      [("", 8, MUTED, False)],
      [("另有 3 种复活方式：看广告（免费）/ 复活令牌 / 花钻石，每种每局限用一次。", 11, DIM, False)]],
     spacing=1.32, space_after=3)

# ══════════════════════════════════════════════
# 10. 成长与竞争
# ══════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); bg(s)
header(s, "03 / 核心功能", "成长与竞争：段位、成就、每日挑战")

text(s, Inches(0.92), Inches(1.9), Inches(5.6), Inches(0.35),
     [[("5 大段位（按最高波次解锁）", 13.5, ACCENT2, True)]])
rows = [("🥉 青铜战士", "默认"), ("🥈 白银骑士", "到达第 5 波"),
        ("🥇 黄金勇士", "到达第 10 波"), ("💎 钻石英雄", "到达第 15 波"),
        ("👑 传说至尊", "到达第 20 波")]
table(s, Inches(0.92), Inches(2.32), Inches(5.4), [2.8, 2.6],
      ["段位", "解锁条件"], rows, row_h=Inches(0.52), font_size=11)

text(s, Inches(6.6), Inches(1.9), Inches(5.8), Inches(0.35),
     [[("长期目标体系", 13.5, ACCENT2, True)]])
text(s, Inches(6.6), Inches(2.35), Inches(5.8), Inches(3.6),
     [[("🏅 12 项成就徽章", 12, TEXT, True)],
      [("累计击杀 / 到达波次 / 连杀 / 单局得分 / 武器收集 / 连续签到六类", 10.5, MUTED, False)],
      [("", 7, MUTED, False)],
      [("🎯 每日挑战", 12, TEXT, True)],
      [("按你的水平自动生成分数目标，达成领金币和钻石", 10.5, MUTED, False)],
      [("", 7, MUTED, False)],
      [("📅 7 天连续签到", 12, TEXT, True)],
      [("第 7 天送大礼包：500 金币 + 5 钻石 + 复活令牌", 10.5, MUTED, False)],
      [("", 7, MUTED, False)],
      [("📜 战绩记录", 12, TEXT, True)],
      [("个人最高分 / 最高波次 / 总击杀，Top 20 历史排行", 10.5, MUTED, False)]],
     spacing=1.3, space_after=2)

box(s, Inches(0.92), Inches(5.35), Inches(5.4), Inches(1.4), fill=CARD2, line=LINE, radius=0.03)
pic_fit(s, ASSETS / "06-rank.png", Inches(1.02), Inches(5.43), Inches(5.2), Inches(1.25))
box(s, Inches(6.6), Inches(5.35), Inches(5.8), Inches(1.4), fill=CARD2, line=LINE, radius=0.03)
pic_fit(s, ASSETS / "04-daily.png", Inches(6.7), Inches(5.43), Inches(5.6), Inches(1.25))
text(s, Inches(0.92), Inches(6.85), Inches(11.5), Inches(0.3),
     [[("左：排行榜（段位 / 战绩 / 成就墙 / 每日挑战）　　右：每日奖励（7 天签到网格）", 10, DIM, False)]])

# ══════════════════════════════════════════════
# 11. 手机适配 + 操作
# ══════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); bg(s)
header(s, "04 / 快速上手", "操作方式与手机适配")

box(s, Inches(0.92), Inches(1.9), Inches(5.6), Inches(2.5), fill=CARD, line=LINE, radius=0.05)
text(s, Inches(1.2), Inches(2.08), Inches(5.0), Inches(0.32), [[("💻 电脑端", 14, TEXT, True)]])
text(s, Inches(1.2), Inches(2.5), Inches(5.0), Inches(1.75),
     [[("移动　W A S D", 11.5, MUTED, False)],
      [("瞄准　鼠标（360° 自由瞄准）", 11.5, MUTED, False)],
      [("射击　鼠标左键（可购买自动射击）", 11.5, MUTED, False)],
      [("闪避　空格（2 秒冷却）", 11.5, MUTED, False)],
      [("暂停　ESC（打开商店 / 升级）", 11.5, MUTED, False)]],
     spacing=1.35, space_after=3)

box(s, Inches(6.82), Inches(1.9), Inches(5.6), Inches(2.5), fill=CARD, line=LINE, radius=0.05)
text(s, Inches(7.1), Inches(2.08), Inches(5.0), Inches(0.32), [[("📱 手机端", 14, TEXT, True)]])
text(s, Inches(7.1), Inches(2.5), Inches(5.0), Inches(1.75),
     [[("移动　左侧虚拟摇杆", 11.5, MUTED, False)],
      [("瞄准　右侧射击键 + 触控瞄准", 11.5, MUTED, False)],
      [("闪避　滑动手势", 11.5, MUTED, False)],
      [("换枪　右下角切换按钮", 11.5, MUTED, False)],
      [("暂停　右上角暂停键", 11.5, MUTED, False)]],
     spacing=1.35, space_after=3)

text(s, Inches(0.92), Inches(4.62), Inches(11.5), Inches(0.35),
     [[("手机不是「能打开」，是「玩得爽」", 13.5, ACCENT2, True)]])
text(s, Inches(0.92), Inches(5.05), Inches(11.5), Inches(1.9),
     [[("· 虚拟摇杆 + 触控瞄准，符合手游操作习惯", 11.5, MUTED, False)],
      [("· 滑动手势闪避，不用在屏幕上找按钮", 11.5, MUTED, False)],
      [("· 竖屏自动提示旋转到横屏，获得最佳体验", 11.5, MUTED, False)],
      [("· 刘海屏安全区适配，全面屏手机上按钮不会被遮挡", 11.5, MUTED, False)],
      [("· 移动端自动降低特效开销，优先保住帧率流畅度", 11.5, MUTED, False)]],
     spacing=1.35, space_after=3)

# ══════════════════════════════════════════════
# 12. 技术复盘
# ══════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); bg(s)
header(s, "05 / 开发与技术复盘", "一个人从零写完的完整游戏",
       "没有用任何游戏引擎、没有任何美术素材——碰撞检测、粒子特效、波次系统、BOSS 战、武器升级树、移动端触控全部手写")

stats = [("8,600+", "行手写代码", "JS + CSS + HTML"),
         ("0", "第三方框架 / 引擎", "纯 Vanilla JS"),
         ("60 fps", "移动端实测帧率", "最重特效皮肤")]
cw = Inches(3.71); gap = Inches(0.19); x0 = Inches(0.92)
for i, (num, lab, sub) in enumerate(stats):
    x = x0 + i * (cw + gap)
    box(s, x, Inches(1.95), cw, Inches(1.4), fill=CARD, line=LINE, radius=0.06)
    box(s, x, Inches(1.95), cw, Inches(0.045), fill=ACCENT)
    text(s, x + Inches(0.26), Inches(2.18), cw - Inches(0.5), Inches(0.5), [[(num, 26, TEXT, True)]])
    text(s, x + Inches(0.26), Inches(2.72), cw - Inches(0.5), Inches(0.3), [[(lab, 11.5, MUTED, False)]])
    text(s, x + Inches(0.26), Inches(3.0), cw - Inches(0.5), Inches(0.28), [[(sub, 10, DIM, False)]])

rows = [("渲染", "Canvas 2D + requestAnimationFrame 自研游戏循环", "零依赖，任何浏览器直接跑"),
        ("框架", "纯 Vanilla JS（无 React / Vue / jQuery）", "游戏状态高频变更，虚拟 DOM 反成负担"),
        ("后端", "Supabase REST API 直连（未装 SDK）", "好友复活需跨设备同步，不想引依赖"),
        ("数据埋点", "客户端事件采集 → 批量上报 → 实时看板", "用真实数据驱动决策"),
        ("皮肤美术", "全部程序化 Canvas 绘制（零图片资源）", "没有美术，用代码反而更系统"),
        ("移动端", "虚拟摇杆 + 触控手势 + 性能自动降级", "手游用户占三分之一，必须真能玩"),
        ("离线可用", "Service Worker 缓存", "网络差时也能加载")]
table(s, Inches(0.92), Inches(3.62), Inches(11.5), [1.7, 5.6, 4.2],
      ["技术点", "实现方式", "为什么这么做"], rows, row_h=Inches(0.42), font_size=10)

# ══════════════════════════════════════════════
# 13. 数据 + 踩坑
# ══════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); bg(s)
header(s, "05 / 开发与技术复盘", "用真实数据说话 & 踩过的坑")

box(s, Inches(0.92), Inches(1.9), Inches(5.6), Inches(4.9), fill=CARD, line=LINE, radius=0.05)
text(s, Inches(1.2), Inches(2.1), Inches(5.0), Inches(0.35),
     [[("做了完整的数据埋点", 14, ACCENT2, True)]])
text(s, Inches(1.2), Inches(2.55), Inches(5.0), Inches(3.9),
     [[("所有关键行为（开局 / 死亡 / 复活 / 广告 / 购买 / 签到 / 成就）都在客户端埋点，落到 Supabase 后在看板实时聚合。", 11, MUTED, False)],
      [("", 8, MUTED, False)],
      [("实测数据：", 11.5, TEXT, True)],
      [("· 累计 1,996 条埋点事件，覆盖 24 个活跃日", 11, MUTED, False)],
      [("· 累计 133 局对局、56 次复活", 11, MUTED, False)],
      [("· 商店购买中 71% 是皮肤", 11, MUTED, False)],
      [("· 移动端事件占 35%", 11, MUTED, False)],
      [("", 8, MUTED, False)],
      [("数据驱动的一个例子：", 11.5, GREEN, True)],
      [("数据显示 92% 的开局都用手枪（免费初始武器），说明多数玩家还没打到能解锁新武器的波次就阵亡了——这直接指向「前期难度偏高 / 武器解锁门槛偏贵」的调优方向。", 11, MUTED, False)]],
     spacing=1.3, space_after=2)

box(s, Inches(6.82), Inches(1.9), Inches(5.6), Inches(4.9), fill=CARD, line=AMBER, radius=0.05)
text(s, Inches(7.1), Inches(2.1), Inches(5.0), Inches(0.35),
     [[("过程中踩过的坑", 14, AMBER, True)]])
text(s, Inches(7.1), Inches(2.55), Inches(5.0), Inches(3.9),
     [[("① 好友复活怎么跨设备？", 11.5, TEXT, True)],
      [("换设备后身份标识不一样，最初救不回来。解法：登录时额外生成一个由账号派生的跨设备标识，复活令牌同时投递给两条身份，任一命中即成功。", 10.5, MUTED, False)],
      [("", 7, MUTED, False)],
      [("② 数据库权限把功能卡死了", 11.5, TEXT, True)],
      [("建表向导默认只给「已登录」角色写权限，但客户端用的是匿名密钥，导致写入被拒。解法：显式给匿名角色授予最小必要权限。", 10.5, MUTED, False)],
      [("", 7, MUTED, False)],
      [("③ 暗色皮肤在暗背景上看不见", 11.5, TEXT, True)],
      [("「暗影刺客」名字要求它暗，但战场本身是近黑色。解法：不靠填充色表达存在感，改用发光描边与光点。", 10.5, MUTED, False)],
      [("", 7, MUTED, False)],
      [("④ 敌人围攻时玩家被淹没", 11.5, TEXT, True)],
      [("实测截图里角色只剩一团红。解法：给玩家加一层深色底衬，保证剪影永远可辨。", 10.5, MUTED, False)]],
     spacing=1.3, space_after=2)

# ══════════════════════════════════════════════
# 14. 项目结构 + 边界
# ══════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); bg(s)
header(s, "05 / 开发与技术复盘", "项目结构")

rows = [("js/game.js", "游戏引擎：实体系统、游戏循环、波次、BOSS、商店、UI", "~4,770 行"),
        ("js/skins.js", "皮肤程序化渲染引擎（6 套皮肤 + 分层绘制）", "~460 行"),
        ("css/style.css", "全部样式 + 响应式 + 触控布局", "~1,940 行"),
        ("js/supabase-db.js", "Supabase REST 封装（好友复活跨设备同步）", "~190 行"),
        ("js/analytics.js", "事件埋点与批量上报", "~190 行"),
        ("dashboard.html", "独立数据看板（实时读取线上数据）", "~430 行")]
table(s, Inches(0.92), Inches(1.95), Inches(11.5), [2.8, 6.9, 1.8],
      ["文件", "内容", "规模"], rows, row_h=Inches(0.6), font_size=11)

box(s, Inches(0.92), Inches(5.85), Inches(11.5), Inches(1.1), fill=CARD, line=AMBER, radius=0.05)
text(s, Inches(1.25), Inches(6.02), Inches(10.9), Inches(0.85),
     [[("诚实的边界：", 12.5, AMBER, True),
       ("皮肤购买目前仍是模拟流程（点确认即到账），未接入真实支付网关；", 11, MUTED, False)],
      [("广告用的是 Monetag 真实广告位；数据看板里的「死亡波次分布」「场均时长」等指标包含开发与自动化测试期间的数据，不作为真实难度曲线引用。", 11, MUTED, False)]],
     spacing=1.3, space_after=2)

# ══════════════════════════════════════════════
# 15. 结尾
# ══════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); bg(s)
box(s, 0, 0, SW, Inches(0.05), fill=ACCENT)
text(s, Inches(1.0), Inches(2.1), Inches(11.3), Inches(1.0),
     [[("不用下载，现在就能玩", 42, TEXT, True)]], align=PP_ALIGN.CENTER)
text(s, Inches(1.0), Inches(3.25), Inches(11.3), Inches(0.5),
     [[("生存竞技场 Survival Arena", 17, ACCENT2, False)]], align=PP_ALIGN.CENTER)
text(s, Inches(1.0), Inches(4.05), Inches(11.3), Inches(1.6),
     [[("🎮 开始游玩：https://survival-arena.com/", 15, GREEN, True)],
      [("📊 数据看板：https://survival-arena.com/dashboard.html", 13, MUTED, False)],
      [("💻 开源代码：https://github.com/AAA-Geeker/Survival-Arena---", 13, MUTED, False)]],
     align=PP_ALIGN.CENTER, spacing=1.6)

out = BASE / "生存竞技场-游戏介绍.pptx"
prs.save(str(out))
tmp = BASE / "_tmp.pptx"
if tmp.exists(): tmp.unlink()
print(f"OK saved: {out}")
print(f"slides: {len(prs.slides._sldIdLst)}")



